"""LangGraph workflow router and business-tool implementations."""

from __future__ import annotations

import csv
import difflib
import io
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, TypedDict

from langgraph.graph import END, START, StateGraph


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "data" / "AI_Agent_Workflow_Assessment.xlsx"
DATA_DIR = PROJECT_ROOT / "data"


@dataclass(frozen=True)
class WorkflowDefinition:
    workflow_id: str
    name: str
    trigger: str
    inputs: str
    steps: str
    decision_logic: str
    tools_required: str
    expected_output: str


class AgentState(TypedDict, total=False):
    request: str
    inputs: dict[str, Any]
    workflow_id: str
    catalog_path: str
    output: dict[str, Any]


def load_workflow_catalog(path: str | Path = DEFAULT_WORKBOOK) -> dict[str, WorkflowDefinition]:
    """Load workflow definitions from the assessment workbook, including appended rows."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Workflows" not in workbook.sheetnames:
        raise ValueError("The workflow workbook must contain a 'Workflows' sheet.")
    rows = workbook["Workflows"].iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    positions = {name: index for index, name in enumerate(headers)}
    required = {
        "Workflow_ID", "Workflow_Name", "Trigger", "Inputs", "Steps",
        "Decision_Logic", "Tools_Required", "Expected_Output",
    }
    missing = required - positions.keys()
    if missing:
        raise ValueError(f"Workflow workbook is missing columns: {', '.join(sorted(missing))}")

    catalog: dict[str, WorkflowDefinition] = {}
    for row in rows:
        values = {key: str(row[index] or "").strip() for key, index in positions.items() if index < len(row)}
        workflow_id = values.get("Workflow_ID", "")
        if not workflow_id:
            continue
        if workflow_id in catalog:
            raise ValueError(f"Duplicate workflow ID in workbook: {workflow_id}")
        catalog[workflow_id] = WorkflowDefinition(
            workflow_id=workflow_id,
            name=values.get("Workflow_Name", ""),
            trigger=values.get("Trigger", ""),
            inputs=values.get("Inputs", ""),
            steps=values.get("Steps", ""),
            decision_logic=values.get("Decision_Logic", ""),
            tools_required=values.get("Tools_Required", ""),
            expected_output=values.get("Expected_Output", ""),
        )
    workbook.close()
    if not catalog:
        raise ValueError("No workflow definitions were found in the workbook.")
    return catalog


ROUTING_PHRASES: dict[str, tuple[str, ...]] = {
    "WF001": ("restock", "restocking", "low stock", "reorder", "need replenishing"),
    "WF002": ("vendor price", "validate product prices", "price differs", "price difference", "price validation"),
    "WF003": ("vendor file", "vendor spreadsheet", "invalid rows", "cleaned dataset", "process this spreadsheet"),
    "WF004": ("seo content", "product description", "seo title", "product copy", "generate content"),
    "WF005": ("order status", "where is order", "tracking number", "shipment status", "order id"),
    "WF006": ("duplicate product", "duplicate products", "likely duplicates", "product duplicates"),
    "WF007": ("campaign brief", "marketing campaign", "campaign for", "campaign goal"),
    "WF008": ("keyword", "search intent", "classify these keywords", "map keywords"),
    "WF009": ("assign a task", "assign this task", "best available", "employee task", "task assignment"),
    "WF010": ("performance report", "workflow logs", "failing most", "failure rate", "execution time"),
}


def select_workflow(request: str, catalog: dict[str, WorkflowDefinition] | None = None) -> tuple[str | None, list[str]]:
    """Route by phrase evidence, returning candidate IDs when clarification is needed."""
    catalog = catalog or load_workflow_catalog()
    query = re.sub(r"\s+", " ", request.lower()).strip()
    query_terms = set(re.findall(r"[a-z0-9]+", query))
    scores: dict[str, int] = {}
    for workflow_id, definition in catalog.items():
        phrases = list(ROUTING_PHRASES.get(workflow_id, ()))
        phrases.extend((definition.name.lower(), definition.trigger.lower()))
        score = sum(2 for phrase in phrases if phrase and phrase in query)
        title_terms = set(re.findall(r"[a-z0-9]+", definition.name.lower()))
        score += len(query_terms & title_terms)
        if score:
            scores[workflow_id] = score
    if not scores:
        return None, []
    highest = max(scores.values())
    matches = sorted(key for key, score in scores.items() if score == highest)
    return (matches[0], []) if len(matches) == 1 else (None, matches)


def _number(value: Any, default: float = 0.0) -> float:
    try:
        return float(str(value).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return default


def _normalized(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def _read_table(inputs: dict[str, Any], key: str, default_name: str, filename_key: str | None = None) -> list[dict[str, Any]]:
    """Read CSV/XLSX, in-memory rows, or the checked-in fixture for a tool input."""
    source = inputs.get(key) or DATA_DIR / default_name
    filename = str(inputs.get(filename_key or f"{key}_filename", "") or source)
    suffix = Path(filename).suffix.lower()
    if isinstance(source, list):
        return [dict(row) for row in source]
    if isinstance(source, dict):
        return [dict(source)]
    if isinstance(source, bytes):
        stream: Any = io.BytesIO(source)
    elif hasattr(source, "getvalue"):
        stream = io.BytesIO(source.getvalue())
    elif isinstance(source, (str, Path)) and Path(source).is_file():
        stream = Path(source)
        suffix = stream.suffix.lower()
    elif isinstance(source, str):
        stream = io.StringIO(source)
        suffix = suffix or ".csv"
    else:
        raise ValueError(f"Could not read '{key}'. Upload CSV/XLSX data or provide rows.")

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            return [
                {headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]}
                for row in rows if any(value is not None for value in row)
            ]
        finally:
            workbook.close()

    if isinstance(stream, Path):
        handle = stream.open("r", encoding="utf-8-sig", newline="")
    elif isinstance(stream, io.BytesIO):
        handle = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    else:
        handle = stream
    try:
        return [dict(row) for row in csv.DictReader(handle)]
    finally:
        handle.close()


def _field(row: dict[str, Any], *names: str, default: Any = "") -> Any:
    normalized = {_normalized(key).replace(" ", "_"): value for key, value in row.items()}
    aliases = {
        "sku": ("sku", "vendor_sku", "product_sku", "item_sku"),
        "product_name": ("product_name", "name", "product", "item_name"),
        "current_stock": ("current_stock", "stock", "quantity", "on_hand"),
        "minimum_stock": ("minimum_stock", "min_stock", "reorder_threshold", "threshold"),
        "price": ("price", "internal_price", "product_price", "our_price"),
        "vendor_price": ("vendor_price", "price", "unit_price"),
        "order_id": ("order_id", "order_number", "id"),
        "customer_email": ("customer_email", "email"),
    }
    candidates: list[str] = []
    for name in names:
        candidates.extend(aliases.get(name, (name,)))
    return next((normalized[name] for name in candidates if name in normalized), default)


def _openai_text(prompt: str, enabled: bool = False) -> tuple[str | None, dict[str, Any]]:
    if not enabled:
        return None, {"service": "OpenAI", "method": "Responses API", "status": "skipped: not enabled"}
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key:
        return None, {"service": "OpenAI", "method": "Responses API", "status": "skipped: no API key"}
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    try:
        from openai import OpenAI

        response = OpenAI(api_key=api_key).responses.create(model=model, input=prompt)
        return response.output_text, {"service": "OpenAI", "method": "Responses API", "model": model, "status": "success"}
    except Exception as exc:
        return None, {"service": "OpenAI", "method": "Responses API", "model": model, "status": f"fallback: {type(exc).__name__}"}


def _product_content(inputs: dict[str, Any]) -> dict[str, Any]:
    fields = {
        "product name": str(inputs.get("product_name", "")).strip(),
        "category": str(inputs.get("category", "")).strip(),
        "attributes": str(inputs.get("attributes", "")).strip(),
        "material": str(inputs.get("material", "")).strip(),
        "color": str(inputs.get("color", "")).strip(),
        "target audience": str(inputs.get("target_audience", "")).strip(),
    }
    if not fields["product name"]:
        return {"status": "needs_input", "message": "Please provide the product name before generating content.", "tools_used": ["text_validator"]}
    missing = [name for name, value in fields.items() if not value and name != "product name"]
    known = [f"{name}: {value}" for name, value in fields.items() if value]
    prompt = ("Write ecommerce product content using only the supplied facts. Do not infer or invent attributes. "
              "Mark missing details explicitly. Return product description, short description, SEO title, and meta description.\n"
              + "\n".join(known) + "\nMissing: " + ", ".join(missing))
    generated, api_call = _openai_text(prompt, enabled=bool(inputs.get("use_openai", False)))
    if generated:
        content: Any = generated
    else:
        name = fields["product name"]
        features = [value for key, value in fields.items() if key not in {"product name", "target audience"} and value]
        details = ", ".join(features) if features else "No additional product attributes were supplied."
        audience = f" Created for {fields['target audience']}." if fields["target audience"] else ""
        missing_note = "Not provided: " + ", ".join(missing) if missing else ""
        content = {
            "product_description": f"{name}. {details}.{audience} {missing_note}".strip(),
            "short_description": f"{name}: {details}.",
            "seo_title": f"{name} | {fields['category']}".strip(" |"),
            "meta_description": f"Discover {name}. {missing_note}".strip(),
            "missing_information": missing,
        }
    return {"status": "completed", "content": content, "tools_used": ["text_validator", "product_content_generator"], "api_calls": [api_call]}


def _campaign_brief(inputs: dict[str, Any]) -> dict[str, Any]:
    goal = str(inputs.get("campaign_goal", inputs.get("goal", ""))).strip()
    start_date = str(inputs.get("start_date", "")).strip()
    end_date = str(inputs.get("end_date", "")).strip()
    missing = [label for label, value in (("campaign goal", goal), ("start date", start_date), ("end date", end_date)) if not value]
    if missing:
        return {"status": "needs_input", "message": "Please provide the " + " and ".join(missing) + " before I generate the campaign brief.",
                "missing_fields": missing, "tools_used": ["campaign_input_validator"]}
    facts = {key: inputs.get(key, "") for key in ("product_list", "target_audience", "promotion", "start_date", "end_date")}
    facts["campaign_goal"] = goal
    generated, api_call = _openai_text(
        "Create a concise campaign brief using only these facts; label unspecified details as not provided.\n" + str(facts),
        enabled=bool(inputs.get("use_openai", False)),
    )
    brief = generated or {
        "objective": goal, "audience": facts["target_audience"] or "Not provided",
        "products": facts["product_list"] or "Not provided", "promotion": facts["promotion"] or "Not provided",
        "messaging": f"Lead with the campaign objective: {goal}. Do not claim unprovided product benefits.",
        "channels": ["Email", "Website", "Social media"], "timeline": {"start": start_date, "end": end_date},
        "checklist": ["Confirm products and promotion", "Approve creative", "Schedule channels", "Review campaign results"],
    }
    return {"status": "completed", "brief": brief, "tools_used": ["campaign_input_validator", "campaign_brief_generator"], "api_calls": [api_call]}


def _run_workflow(workflow_id: str, inputs: dict[str, Any]) -> dict[str, Any]:
    if workflow_id == "WF001":
        rows = _read_table(inputs, "inventory_file", "inventory.csv")
        items = []
        for row in rows:
            stock = _number(_field(row, "current_stock"))
            threshold = _number(inputs.get("minimum_stock", _field(row, "minimum_stock")))
            if stock < threshold:
                items.append({"sku": _field(row, "sku"), "product_name": _field(row, "product_name"),
                              "current_stock": stock, "minimum_stock": threshold,
                              "suggested_reorder_quantity": max(0, int(threshold - stock))})
        return {"status": "completed", "items_to_restock": items, "count": len(items), "tools_used": ["csv_reader", "stock_calculator"]}

    if workflow_id == "WF002":
        products = _read_table(inputs, "products_file", "products.csv")
        vendors = _read_table(inputs, "vendor_prices_file", "vendor_prices.csv")
        vendor_by_sku = {str(_field(row, "sku")).strip().lower(): row for row in vendors if _field(row, "sku")}
        report = []
        for product in products:
            sku = str(_field(product, "sku")).strip()
            vendor = vendor_by_sku.get(sku.lower())
            if not vendor:
                report.append({"sku": sku, "product_name": _field(product, "product_name"), "status": "vendor price missing"})
                continue
            internal, vendor_price = _number(_field(product, "price")), _number(_field(vendor, "vendor_price"))
            difference = round(abs(vendor_price - internal) / abs(internal) * 100, 2) if internal else None
            report.append({"sku": sku, "product_name": _field(product, "product_name"), "internal_price": internal,
                           "vendor_price": vendor_price, "difference_percent": difference,
                           "flagged": difference is not None and difference > _number(inputs.get("tolerance_percent", 10))})
        return {"status": "completed", "price_report": report, "exception_count": sum(bool(row.get("flagged")) for row in report),
                "tools_used": ["csv_reader", "sku_matcher", "price_calculator"]}

    if workflow_id == "WF003":
        rows = _read_table(inputs, "vendor_file", "vendor_upload.csv", "vendor_filename")
        cleaned, invalid = [], []
        for number, row in enumerate(rows, start=2):
            normalized = {_normalized(key).replace(" ", "_"): value for key, value in row.items()}
            sku, name = _field(row, "sku"), _field(row, "product_name")
            reasons = [label for label, value in (("missing SKU", sku), ("missing product name", name)) if not str(value).strip()]
            canonical = dict(normalized, sku=str(sku).strip(), product_name=str(name).strip())
            if reasons:
                invalid.append({"row_number": number, "reasons": reasons, "data": canonical})
            else:
                cleaned.append(canonical)
        return {"status": "completed", "cleaned_rows": cleaned, "invalid_rows": invalid,
                "summary": {"received": len(rows), "valid": len(cleaned), "invalid": len(invalid)},
                "tools_used": ["csv_xlsx_reader", "column_normalizer", "required_field_validator"]}

    if workflow_id == "WF004":
        return _product_content(inputs)

    if workflow_id == "WF005":
        request_order = re.search(r"\bORD[- ]?\d+\b", str(inputs.get("request", "")), re.IGNORECASE)
        identifier = str(inputs.get("order_id", "") or (request_order.group(0) if request_order else "")).strip().lower()
        email = str(inputs.get("customer_email", "")).strip().lower()
        if not identifier and not email:
            return {"status": "needs_input", "message": "Please provide an order ID or customer email.", "tools_used": ["identifier_validator"]}
        orders = _read_table(inputs, "orders_file", "orders.csv")
        found = next((row for row in orders if (identifier and str(_field(row, "order_id")).lower() == identifier)
                      or (email and str(_field(row, "customer_email")).lower() == email)), None)
        if not found:
            return {"status": "not_found", "message": "No order matched that identifier. Please check the order ID or provide the customer email.",
                    "tools_used": ["order_lookup"]}
        return {"status": "completed", "order": found, "tools_used": ["order_lookup", "shipment_lookup"]}

    if workflow_id == "WF006":
        rows = _read_table(inputs, "catalog_file", "product_catalog.csv")
        groups, used = [], set()
        for left_index, left in enumerate(rows):
            left_sku, left_name = str(_field(left, "sku")).strip().lower(), _normalized(_field(left, "product_name"))
            for right_index in range(left_index + 1, len(rows)):
                if right_index in used:
                    continue
                right = rows[right_index]
                right_sku, right_name = str(_field(right, "sku")).strip().lower(), _normalized(_field(right, "product_name"))
                same_sku = bool(left_sku and left_sku == right_sku)
                similarity = difflib.SequenceMatcher(None, left_name, right_name).ratio()
                if same_sku or similarity >= _number(inputs.get("similarity_threshold", 0.84)):
                    groups.append({"products": [left, right], "confidence": "definite" if same_sku else "possible",
                                   "matching_fields": ["sku"] if same_sku else ["product_name"], "name_similarity": round(similarity, 3)})
                    used.add(right_index)
        return {"status": "completed", "duplicate_groups": groups, "count": len(groups), "tools_used": ["catalog_reader", "name_similarity"]}

    if workflow_id == "WF007":
        return _campaign_brief(inputs)

    if workflow_id == "WF008":
        rows = _read_table(inputs, "keywords_file", "keywords.csv")
        seen, report = set(), []
        for row in rows:
            keyword = str(row.get("keyword", row.get("Keyword", ""))).strip()
            normalized = _normalized(keyword)
            if not keyword or normalized in seen:
                continue
            seen.add(normalized)
            words = set(normalized.split())
            if words & {"buy", "order", "coupon", "discount", "purchase"}:
                intent = "transactional"
            elif words & {"best", "review", "compare", "vs", "top"}:
                intent = "commercial"
            elif words & {"what", "how", "why", "when", "guide", "meaning"}:
                intent = "informational"
            elif words & {"login", "official", "website", "brand"}:
                intent = "navigational"
            else:
                intent = "commercial" if len(words) > 1 else "informational"
            priority = "high" if intent == "transactional" else "medium" if intent in {"commercial", "navigational"} else "low"
            category = row.get("category", row.get("Category", "")) or inputs.get("category", "General")
            page = "Product page" if intent == "transactional" else "Comparison/category page" if intent == "commercial" else "Guide/article" if intent == "informational" else "Brand page"
            report.append({"keyword": keyword, "intent": intent, "category": category, "priority": priority, "recommended_target_page": page})
        return {"status": "completed", "keyword_report": report, "count": len(report),
                "tools_used": ["keyword_reader", "duplicate_filter", "intent_classifier", "page_mapper"]}

    if workflow_id == "WF009":
        request = str(inputs.get("request", ""))
        task = str(inputs.get("task_description", "") or request).strip()
        employees = _read_table(inputs, "employees_file", "employees.csv")
        required_skills = inputs.get("required_skills", [])
        if isinstance(required_skills, str):
            required_skills = [skill.strip().lower() for skill in re.split(r"[,;]", required_skills) if skill.strip()]
        if not required_skills:
            terms = set(_normalized(task).split())
            required_skills = [skill for skill in ("python", "javascript", "sql", "design", "frontend", "backend", "testing", "data", "developer") if skill in terms]
        ranked = []
        for employee in employees:
            name = employee.get("employee", employee.get("name", ""))
            skills = [skill.strip().lower() for skill in re.split(r"[,;|]", str(employee.get("skills", ""))) if skill.strip()]
            workload = _number(employee.get("current_workload", employee.get("workload", 0)))
            capacity = _number(employee.get("capacity", 100), 100)
            available = capacity - workload
            matched = sorted(set(required_skills) & set(skills))
            if available > 0 and (not required_skills or matched):
                ranked.append({"employee": name, "matched_skills": matched, "available_capacity": available,
                               "score": len(matched) * 100 + available / max(capacity, 1) * 10})
        ranked.sort(key=lambda candidate: candidate["score"], reverse=True)
        if not ranked:
            return {"status": "escalated", "message": "No employee with the required skills and available capacity was found. Please review or expand the candidate list.",
                    "required_skills": required_skills, "tools_used": ["employee_reader", "skill_capacity_ranker"]}
        selected = ranked[0]
        priority = str(inputs.get("priority", "normal"))
        if priority == "normal" and re.search(r"\burgent\b", request, re.IGNORECASE):
            priority = "urgent"
        return {"status": "completed", "assignment": {
            "employee": selected["employee"],
            "reasoning": f"Matched skills: {', '.join(selected['matched_skills']) or 'No specific skills required'}; available capacity: {selected['available_capacity']}.",
            "priority": priority, "deadline": inputs.get("deadline", "Not provided"), "task_summary": task,
        }, "ranked_candidates": ranked, "tools_used": ["employee_reader", "skill_capacity_ranker"]}

    if workflow_id == "WF010":
        rows = _read_table(inputs, "logs_file", "workflow_logs.csv")
        total = len(rows)
        failures = [row for row in rows if str(row.get("status", "")).lower() in {"failed", "failure", "error"}]
        average = round(sum(_number(row.get("duration_seconds")) for row in rows) / total, 2) if total else 0
        failure_rate = round(len(failures) / total * 100, 2) if total else 0
        errors: dict[str, int] = {}
        for row in failures:
            message = str(row.get("error_message", "Unspecified error")).strip() or "Unspecified error"
            errors[message] = errors.get(message, 0) + 1
        slow_threshold = _number(inputs.get("slow_step_threshold_seconds", 30), 30)
        slow_steps = [{"workflow_id": row.get("workflow_id"), "step": row.get("step"), "duration_seconds": _number(row.get("duration_seconds"))}
                      for row in rows if _number(row.get("duration_seconds")) > slow_threshold]
        flagged = failure_rate > _number(inputs.get("failure_threshold_percent", 10), 10) or average > slow_threshold
        recommendations = []
        if failure_rate > 10:
            recommendations.append("Investigate recurring failures and add targeted retries or input validation.")
        if slow_steps:
            recommendations.append("Profile the slow steps and reduce repeated I/O or external calls.")
        if not recommendations:
            recommendations.append("Continue monitoring; no threshold breach was found in this sample.")
        return {"status": "completed", "metrics": {"runs": total, "failures": len(failures), "failure_rate_percent": failure_rate,
                "average_execution_seconds": average, "threshold_breached": flagged},
                "frequent_errors": sorted(({"error": error, "count": count} for error, count in errors.items()), key=lambda item: item["count"], reverse=True),
                "slow_steps": slow_steps, "recommendations": recommendations,
                "tools_used": ["execution_log_reader", "performance_calculator", "report_generator"]}

    definition = load_workflow_catalog(inputs.get("catalog_path", DEFAULT_WORKBOOK)).get(workflow_id)
    return {"status": "plan_only", "workflow": definition.name if definition else workflow_id,
            "trigger": definition.trigger if definition else "", "required_inputs": definition.inputs if definition else "",
            "execution_steps": definition.steps if definition else "", "decision_logic": definition.decision_logic if definition else "",
            "message": "Workflow is registered. Add a specialized handler to automate its business actions.",
            "tools_used": ["workflow_catalog", "generic_plan_handler"]}


def build_agent(catalog_path: str | Path = DEFAULT_WORKBOOK):
    """Build the LangGraph request-routing and workflow-execution graph."""
    catalog_path = str(catalog_path)

    def route(state: AgentState) -> dict[str, Any]:
        catalog = load_workflow_catalog(state.get("catalog_path", catalog_path))
        explicit_id = state.get("workflow_id")
        if explicit_id:
            if explicit_id not in catalog:
                return {"output": {"status": "not_found", "message": f"Workflow {explicit_id} is not in the workbook catalog.", "tools_used": ["workflow_catalog"]}}
            return {"workflow_id": explicit_id}
        workflow_id, candidates = select_workflow(state.get("request", ""), catalog)
        if workflow_id:
            return {"workflow_id": workflow_id}
        if candidates:
            return {"output": {"status": "needs_clarification", "message": "This request could match more than one workflow.",
                                "candidates": candidates, "tools_used": ["workflow_router"]}}
        return {"output": {"status": "needs_clarification", "message": "I could not match that request to a workflow. Please select one or describe the task more specifically.",
                            "candidates": list(catalog), "tools_used": ["workflow_router"]}}

    def execute(state: AgentState) -> dict[str, Any]:
        workflow_id = state["workflow_id"]
        active_catalog_path = state.get("catalog_path", catalog_path)
        catalog = load_workflow_catalog(active_catalog_path)
        inputs = {**state.get("inputs", {}), "catalog_path": active_catalog_path, "request": state.get("request", "")}
        result = _run_workflow(workflow_id, inputs)
        return {"output": {"workflow_id": workflow_id, "workflow_name": catalog[workflow_id].name, **result}}

    graph = StateGraph(AgentState)
    graph.add_node("route_request", route)
    graph.add_node("execute_workflow", execute)
    graph.add_edge(START, "route_request")
    graph.add_conditional_edges("route_request", lambda state: "execute" if state.get("workflow_id") else "stop",
                                {"execute": "execute_workflow", "stop": END})
    graph.add_edge("execute_workflow", END)
    return graph.compile()


def run_agent(request: str, inputs: dict[str, Any] | None = None, workflow_id: str | None = None,
              catalog_path: str | Path = DEFAULT_WORKBOOK) -> dict[str, Any]:
    result = build_agent(catalog_path).invoke({"request": request, "inputs": inputs or {},
                                               "workflow_id": workflow_id or "", "catalog_path": str(catalog_path)})
    return result.get("output", {"status": "error", "message": "Workflow execution produced no result."})