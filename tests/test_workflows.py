from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from workflow_agent.engine import DEFAULT_WORKBOOK, load_workflow_catalog, run_agent, select_workflow


class WorkflowAgentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = load_workflow_catalog()

    def execute(self, workflow_id, inputs=None, request="run workflow"):
        return run_agent(request, inputs or {}, workflow_id)

    def test_workbook_has_ten_workflows_and_ten_test_prompts(self):
        workbook = load_workbook(DEFAULT_WORKBOOK, read_only=True, data_only=True)
        self.assertEqual(len(self.catalog), 10)
        self.assertEqual(workbook["Test_Questions"].max_row - 1, 10)

    def test_provided_test_requests_route_to_all_ten_workflows(self):
        workbook = load_workbook(DEFAULT_WORKBOOK, read_only=True, data_only=True)
        for workflow_id, request, _ in list(workbook["Test_Questions"].iter_rows(min_row=2, values_only=True)):
            with self.subTest(workflow_id=workflow_id, request=request):
                self.assertEqual(select_workflow(request, self.catalog)[0], workflow_id)

    def test_wf001_restock_threshold_is_strictly_less_than(self):
        result = self.execute("WF001")
        self.assertEqual(result["count"], 2)
        self.assertEqual(result["items_to_restock"][0]["suggested_reorder_quantity"], 8)

    def test_csv_upload_bytes_are_decoded_using_the_uploaded_filename(self):
        result = self.execute("WF001", {
            "inventory_file": b"sku,product_name,current_stock,minimum_stock\nA-1,Sample,1,5\n",
            "inventory_file_filename": "inventory.csv",
        })
        self.assertEqual(result["items_to_restock"][0]["sku"], "A-1")

    def test_wf002_flags_price_differences_over_ten_percent(self):
        result = self.execute("WF002")
        self.assertEqual(result["exception_count"], 2)

    def test_wf002_does_not_flag_a_difference_of_exactly_ten_percent(self):
        result = self.execute("WF002", {
            "products_file": [{"sku": "A-1", "product_name": "Sample", "price": 100}],
            "vendor_prices_file": [{"sku": "A-1", "vendor_price": 110}],
        })
        self.assertEqual(result["price_report"][0]["difference_percent"], 10.0)
        self.assertFalse(result["price_report"][0]["flagged"])

    def test_wf003_reports_invalid_rows_and_cleans_valid_rows(self):
        result = self.execute("WF003")
        self.assertEqual(result["summary"], {"received": 4, "valid": 2, "invalid": 2})

    def test_wf004_marks_missing_attributes_without_inventing_them(self):
        result = self.execute("WF004", {"product_name": "Harbor Mug", "material": "ceramic"})
        self.assertEqual(result["status"], "completed")
        self.assertIn("Not provided: category", result["content"]["product_description"])
        self.assertIn("color", result["content"]["missing_information"])

    def test_openai_is_not_called_without_explicit_opt_in(self):
        with patch.dict("os.environ", {"OPENAI_API_KEY": "test-key"}):
            result = self.execute("WF004", {"product_name": "Harbor Mug"})
        self.assertEqual(result["api_calls"][0]["status"], "skipped: not enabled")

    def test_wf005_finds_order_and_handles_unknown_identifier(self):
        found = self.execute("WF005", {"order_id": "ORD-1001"})
        missing = self.execute("WF005", {"order_id": "ORD-9999"})
        self.assertEqual(found["order"]["tracking_number"], "1Z999AA10123456784")
        self.assertEqual(missing["status"], "not_found")

    def test_wf005_extracts_order_id_from_the_workbook_example_request(self):
        result = run_agent("Where is order ORD-1001?")
        self.assertEqual(result["order"]["order_id"], "ORD-1001")

    def test_wf006_detects_exact_sku_and_possible_name_duplicates(self):
        result = self.execute("WF006")
        self.assertTrue(any(group["confidence"] == "definite" for group in result["duplicate_groups"]))
        self.assertTrue(any(group["confidence"] == "possible" for group in result["duplicate_groups"]))

    def test_wf007_requests_goal_and_dates_before_generation(self):
        result = self.execute("WF007")
        self.assertEqual(result["status"], "needs_input")
        self.assertEqual(set(result["missing_fields"]), {"campaign goal", "start date", "end date"})

    def test_wf007_generates_brief_when_required_fields_are_present(self):
        result = self.execute("WF007", {"campaign_goal": "Launch spring collection", "start_date": "2026-03-01", "end_date": "2026-03-31"})
        self.assertEqual(result["status"], "completed")

    def test_wf008_classifies_intent_and_deduplicates_keywords(self):
        result = self.execute("WF008")
        self.assertEqual(result["count"], 5)
        intents = {row["intent"] for row in result["keyword_report"]}
        self.assertTrue({"informational", "commercial", "transactional", "navigational"}.issubset(intents))

    def test_wf009_selects_available_candidate_with_required_skills(self):
        result = self.execute("WF009", {"task_description": "Build the Python backend", "required_skills": ["python", "backend"], "priority": "urgent"})
        self.assertEqual(result["assignment"]["employee"], "Maya Chen")

    def test_wf009_uses_the_workbook_example_request_for_task_and_urgency(self):
        result = run_agent("Assign this urgent task to the best available developer.")
        self.assertEqual(result["assignment"]["employee"], "Maya Chen")
        self.assertEqual(result["assignment"]["priority"], "urgent")

    def test_wf009_escalates_when_no_suitable_employee_exists(self):
        result = self.execute("WF009", {"required_skills": ["quantum"]})
        self.assertEqual(result["status"], "escalated")

    def test_wf010_calculates_metrics_and_flags_slow_steps(self):
        result = self.execute("WF010")
        self.assertEqual(result["metrics"]["runs"], 10)
        self.assertEqual(result["metrics"]["failure_rate_percent"], 20.0)
        self.assertTrue(result["metrics"]["threshold_breached"])

    def test_new_workbook_row_uses_generic_plan_handler(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "workflows.xlsx"
            shutil.copy(DEFAULT_WORKBOOK, workbook_path)
            workbook = load_workbook(workbook_path)
            workbook["Workflows"].append([
                "WF011", "Returns Eligibility Check", "User asks if a return is eligible",
                "Order ID; purchase date", "Load order; check return window", "If within 30 days, eligible",
                "Order lookup; date calculator", "Eligibility decision and reason",
            ])
            workbook.save(workbook_path)
            workbook.close()
            result = run_agent("Is this return eligible?", {}, "WF011", workbook_path)
            self.assertEqual(result["status"], "plan_only")
            self.assertEqual(result["workflow_name"], "Returns Eligibility Check")


if __name__ == "__main__":
    unittest.main()